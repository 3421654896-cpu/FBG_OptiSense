import os
import threading
import time
import unittest
from pathlib import Path
from unittest import mock

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
from PyQt5 import QtWidgets

import app_JDSU
import fbg_unified_app


HERE = Path(__file__).resolve().parent


class APWorkerExcelParserTests(unittest.TestCase):
    def test_wavelength_meter_identity_accepts_current_and_legacy_units(self):
        self.assertTrue(
            app_JDSU.is_supported_aq6150_identity(
                "YOKOGAWA,AQ6150,91P102177,01.03"
            )
        )
        self.assertTrue(
            app_JDSU.is_supported_aq6150_identity(
                "YOKOGAWA,AQ6150B,9027C2596,01.03"
            )
        )
        self.assertFalse(
            app_JDSU.is_supported_aq6150_identity(
                "YOKOGAWA,AQ6150,UNAPPROVED,01.03"
            )
        )

    def make_worker(self, headers):
        worker = app_JDSU.APWorker("")
        worker.configure_excel_columns(headers)
        self.addCleanup(worker.deleteLater)
        return worker

    def test_real_high_power_table_uses_code_columns_not_power_or_current(self):
        path = HERE / "高功率应力扫描_最终68点标定表.xlsx"
        workbook = load_workbook(path, read_only=True, data_only=True)
        self.addCleanup(workbook.close)
        sheet = workbook.active
        headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        row = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))

        worker = self.make_worker(headers)
        self.assertEqual(worker.excel_code_columns, (14, 15, 16, 17, 18))
        command, _info = worker.excel_operate(iter((row,)))
        expected_codes = (52428, 50687, 8630, 4915, 22118)
        self.assertEqual(
            command,
            app_JDSU.build_single_value_dac_command(expected_codes),
        )
        # These are fractional power/error/current values and must never be
        # selected merely because they occupied the former row[6:11] slice.
        self.assertNotEqual(tuple(row[6:11]), expected_codes)

    def test_legacy_code_headers_are_supported(self):
        headers = (
            "target_nm",
            "code_GAIN",
            "code_SOA",
            "code_PHASE",
            "code_wavelengthA",
            "code_wavelengthB",
        )
        worker = self.make_worker(headers)
        self.assertEqual(worker.excel_code_columns, (1, 2, 3, 4, 5))
        row = (1525.0, 100, 200, 300, 400, 500)
        command, _info = worker.excel_operate(iter((row,)))
        self.assertEqual(
            command,
            app_JDSU.build_single_value_dac_command(row[1:6]),
        )

    def test_missing_code_column_does_not_fall_back_to_current_columns(self):
        headers = (
            "GAIN(mA)",
            "SOA(mA)",
            "PHASE(mA)",
            "WAVE_A(mA)",
            "WAVE_B(mA)",
            "GAIN码",
            "SOA码",
            "PHASE码",
            "WAVE_A码",
        )
        with self.assertRaisesRegex(ValueError, "缺少.*WAVE_B"):
            app_JDSU.resolve_pi11210_excel_code_columns(headers)

    def test_duplicate_channel_code_header_is_rejected_as_ambiguous(self):
        headers = (
            "GAIN码",
            "code_GAIN",
            "SOA码",
            "PHASE码",
            "WAVE_A码",
            "WAVE_B码",
        )
        with self.assertRaisesRegex(ValueError, "歧义.*GAIN"):
            app_JDSU.resolve_pi11210_excel_code_columns(headers)

    def test_fractional_code_is_rejected(self):
        worker = self.make_worker(
            ("GAIN码", "SOA码", "PHASE码", "WAVE_A码", "WAVE_B码")
        )
        with self.assertRaisesRegex(ValueError, "必须是十进制整数"):
            worker.excel_operate(iter(((100, 200, 3.5, 400, 500),)))

    def test_out_of_range_code_is_rejected(self):
        worker = self.make_worker(
            ("GAIN码", "SOA码", "PHASE码", "WAVE_A码", "WAVE_B码")
        )
        with self.assertRaisesRegex(ValueError, "0～58981"):
            worker.excel_operate(iter(((58983, 200, 300, 400, 500),)))

    def test_boolean_code_is_rejected(self):
        worker = self.make_worker(
            ("GAIN码", "SOA码", "PHASE码", "WAVE_A码", "WAVE_B码")
        )
        with self.assertRaisesRegex(ValueError, "必须是十进制整数"):
            worker.excel_operate(iter(((True, 200, 300, 400, 500),)))

    def test_corrupted_suffix_is_accepted_only_after_complete_current_block(self):
        headers = (
            "GAIN(mA)",
            "SOA(mA)",
            "PHASE(mA)",
            "WAVE_A(mA)",
            "WAVE_B(mA)",
            "GAIN��",
            "SOA��",
            "PHASE��",
            "WAVE_A��",
            "WAVE_B��",
        )
        self.assertEqual(
            app_JDSU.resolve_pi11210_excel_code_columns(headers),
            (5, 6, 7, 8, 9),
        )

    def test_bare_channel_names_are_not_a_positional_fallback(self):
        headers = (
            "GAIN(mA)",
            "SOA(mA)",
            "PHASE(mA)",
            "WAVE_A(mA)",
            "WAVE_B(mA)",
            "GAIN",
            "SOA",
            "PHASE",
            "WAVE_A",
            "WAVE_B",
        )
        with self.assertRaisesRegex(ValueError, "未找到完整"):
            app_JDSU.resolve_pi11210_excel_code_columns(headers)


class APWorkerTransportTests(unittest.TestCase):
    @staticmethod
    def manual_frame(frame_type, value=0):
        frame = bytearray(app_JDSU.SINGLE_VALUE_FRAME_SIZE)
        frame[0:4] = bytes((0xFF, 0xFF, 0x01, frame_type))
        frame[4] = value
        return bytes(frame)

    def make_worker(self):
        worker = app_JDSU.APWorker("")
        self.addCleanup(worker.deleteLater)
        return worker

    def test_bulk_parser_reaches_ack_behind_temperature_flood(self):
        temperature = self.manual_frame(0x01, 28)
        ack = self.manual_frame(0x00, app_JDSU.ACK_VALUE)
        buffer = bytearray(temperature * 3900 + ack)

        frames = app_JDSU.extract_manual_scan_frames(buffer)

        self.assertEqual(len(frames), 3901)
        self.assertEqual(app_JDSU.manual_scan_ack_status(frames[-1]), 0x21)
        self.assertEqual(buffer, b"")

    def test_explicit_pi11210_failure_is_not_silently_retried(self):
        worker = self.make_worker()
        worker.flag_queue.put(
            self.manual_frame(0x00, app_JDSU.ACK_ERROR_VALUE)
        )

        with mock.patch.object(app_JDSU, "serial_write") as write:
            with self.assertRaisesRegex(RuntimeError, "ACK=0xE1"):
                worker._send_command_with_ack(b"command")

        write.assert_called_once_with(b"command")

    def test_missing_ack_has_bounded_retry_count(self):
        worker = self.make_worker()
        with mock.patch.object(app_JDSU, "ACK_ATTEMPT_TIMEOUT_S", 0.01), \
                mock.patch.object(app_JDSU, "ACK_WAIT_SLICE_S", 0.002), \
                mock.patch.object(app_JDSU, "serial_write") as write:
            with self.assertRaisesRegex(RuntimeError, "仍未收到"):
                worker._send_command_with_ack(b"command")

        self.assertEqual(write.call_count, app_JDSU.ACK_MAX_ATTEMPTS)

    def test_stop_wakes_ack_wait_without_another_serial_write(self):
        worker = self.make_worker()
        result = []
        with mock.patch.object(app_JDSU, "ACK_ATTEMPT_TIMEOUT_S", 5.0), \
                mock.patch.object(app_JDSU, "serial_write") as write:
            waiter = threading.Thread(
                target=lambda: result.append(
                    worker._send_command_with_ack(b"command")
                )
            )
            waiter.start()
            time.sleep(0.02)
            worker.stop()
            waiter.join(timeout=0.3)

        self.assertFalse(waiter.is_alive())
        self.assertEqual(result, [None])
        write.assert_called_once_with(b"command")

    def test_safe_shutdown_switches_to_extra_and_verifies_shutter(self):
        class FakeSerial:
            is_open = True

            def __init__(self):
                self.buffer = bytearray()

            @property
            def in_waiting(self):
                return len(self.buffer)

            def flush(self):
                pass

            def reset_input_buffer(self):
                self.buffer.clear()

            def read(self, size):
                data = bytes(self.buffer[:size])
                del self.buffer[:size]
                return data

        fake_serial = FakeSerial()
        writes = []
        ack = bytearray(app_JDSU.SINGLE_VALUE_FRAME_SIZE)
        ack[0:4] = b"\xFF\xFF\x02\x00"
        ack[14] = app_JDSU.PI11210_SOA_SHUTTER_MODE
        ack[15] = 0x80 | 0x01 | 0x02 | 0x04

        def capture_write(payload):
            writes.append(payload)
            if payload[0:4] == b"\xFF\xFF\x00\x01":
                fake_serial.buffer.extend(ack)

        worker = self.make_worker()
        with mock.patch.object(app_JDSU, "ser", fake_serial), \
                mock.patch.object(app_JDSU, "serial_write", side_effect=capture_write), \
                mock.patch.object(app_JDSU, "AP_MODE_SWITCH_SETTLE_S", 0.0):
            self.assertTrue(worker._safe_shutdown_laser())

        self.assertEqual(len(writes), 2)
        self.assertEqual(writes[0][0:4], b"\xFF\xFF\x01\x02")
        self.assertEqual(writes[0][8], 2)
        self.assertEqual(writes[1][0:4], b"\xFF\xFF\x00\x01")
        self.assertEqual(writes[1][4:14], b"\x00" * 10)
        self.assertEqual(writes[1][14], app_JDSU.PI11210_SOA_SHUTTER_MODE)


class APWindowLifecycleTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])

    def test_close_button_requests_stop_without_blocking_gui_thread(self):
        class RunningWorker:
            stopped = False

            @staticmethod
            def isRunning():
                return True

            def stop(self):
                self.stopped = True

        window = app_JDSU.ap6150bWindow()
        self.addCleanup(window.deleteLater)
        worker = RunningWorker()
        window.worker = worker

        started = time.monotonic()
        window.ap_thread()
        elapsed = time.monotonic() - started

        self.assertTrue(worker.stopped)
        self.assertLess(elapsed, 0.1)
        self.assertFalse(window.com_btn.isEnabled())
        self.assertEqual(window.com_btn.text(), "正在关闭…")

    def test_standalone_window_ignores_close_until_sweep_finishes(self):
        class Signal:
            callback = None

            def connect(self, callback):
                self.callback = callback

        class Worker:
            def __init__(self):
                self.finished = Signal()
                self.stopped = False

            @staticmethod
            def isRunning():
                return True

            def stop(self):
                self.stopped = True

        class Event:
            ignored = False

            def ignore(self):
                self.ignored = True

        worker = Worker()
        page = type("Page", (), {"worker": worker})()
        owner = type(
            "Owner",
            (),
            {
                "page_ap6150": page,
                "_ap_close_pending": False,
                "_retry_close_after_ap_worker": lambda self: None,
            },
        )()
        event = Event()

        app_JDSU.MainWindow.closeEvent(owner, event)

        self.assertTrue(worker.stopped)
        self.assertTrue(event.ignored)
        self.assertTrue(owner._ap_close_pending)
        self.assertIsNotNone(worker.finished.callback)

    def test_unified_window_also_defers_close_until_sweep_finishes(self):
        class Signal:
            callback = None

            def connect(self, callback):
                self.callback = callback

        class Worker:
            def __init__(self):
                self.finished = Signal()
                self.stopped = False

            @staticmethod
            def isRunning():
                return True

            def stop(self):
                self.stopped = True

        class Event:
            ignored = False

            def ignore(self):
                self.ignored = True

        class StatusBar:
            message = ""

            def showMessage(self, message):
                self.message = message

        worker = Worker()
        status = StatusBar()
        controller = type(
            "Controller",
            (),
            {"page_ap6150": type("Page", (), {"worker": worker})()},
        )()
        owner = type(
            "Owner",
            (),
            {
                "_closing": False,
                "_sweep_close_pending": False,
                "local_controller": controller,
                "_retry_close_after_sweep": lambda self: None,
                "statusBar": lambda self: status,
            },
        )()
        event = Event()

        fbg_unified_app.UnifiedFbgStudio.closeEvent(owner, event)

        self.assertTrue(worker.stopped)
        self.assertTrue(event.ignored)
        self.assertTrue(owner._sweep_close_pending)
        self.assertIn("SOA", status.message)


if __name__ == "__main__":
    unittest.main()
